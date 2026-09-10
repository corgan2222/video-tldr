"""Step 7, the documents: the HTML a browser prints to PDF, the Word file
python-docx builds from the same data as the Markdown, and the PNG the
browser draws from a Mermaid diagram.

No pandoc on purpose: it is not installed, and the note's data is a dict
the renderers read directly, so nothing has to parse Markdown back. The
Mermaid script ships in `assets/` next to this file; nothing is fetched.
"""

from __future__ import annotations

import os
import shutil
import subprocess
from html import escape, unescape
from pathlib import Path

from .analyze import stamp
from .fetch import FetchError, ffmpeg

MERMAID_JS = Path(__file__).parent / "assets" / "mermaid.min.js"
# The diagram is drawn in a window this large at twice the pixel density
# and cut to its content plus this margin, in pixels of the result.
DIAGRAM_WINDOW = (1400, 1000)
DIAGRAM_MARGIN = 24

# Where Chrome and Edge install on Windows; `browser` in config.json
# overrides the search.
BROWSERS = [
    "C:/Program Files/Google/Chrome/Application/chrome.exe",
    "C:/Program Files (x86)/Google/Chrome/Application/chrome.exe",
    "~/AppData/Local/Google/Chrome/Application/chrome.exe",
    "C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe",
    "C:/Program Files/Microsoft/Edge/Application/msedge.exe",
]
PRINT_TIMEOUT_SECONDS = 120

STYLE = """
body { font-family: Segoe UI, Helvetica, Arial, sans-serif; max-width: 52em;
       margin: 2em auto; line-height: 1.45; color: #222; }
img { max-width: 100%; height: auto; page-break-inside: avoid; }
h2 { page-break-after: avoid; }
code { font-family: Consolas, monospace; font-size: 0.95em; }
pre { background: #f6f6f6; border: 1px solid #ddd; border-radius: 4px;
      padding: 0.6em 0.8em; overflow-x: auto; page-break-inside: avoid; }
pre code { font-size: 0.9em; background: none; border: none; }
"""


# What a browser runs instead of opening, once it stands in an `href` or
# a `src`.
RUNNABLE_SCHEMES = ("javascript:", "data:", "vbscript:")
# The one `data:` a note may carry: a picture is not code. A README's
# base64 badge can reach the note inside the model's answer about that
# README, and behind `unsafe:` it is a source no browser loads. SVG stays
# blocked, because an SVG file is a document that can hold a `script`
# element and outside references; a raster format has nowhere to put one.
PICTURE_DATA = "data:image/"
SCRIPTABLE_DATA = "data:image/svg"


def runnable(target: str) -> bool:
    """Whether a browser would run this link or picture target rather than
    open it. Read close to the way a browser reads a URL: entities
    decoded, then every ASCII space and control character dropped, then
    the scheme compared without case. Dropping every space is wider than
    a URL parser goes, which keeps one inside a word, and wider is the
    side to err on here."""
    plain = "".join(c for c in unescape(target) if c > " " and c != "\x7f").lower()
    if plain.startswith(PICTURE_DATA) and not plain.startswith(SCRIPTABLE_DATA):
        return False
    return plain.startswith(RUNNABLE_SCHEMES)


def safe_target(target: str) -> str:
    """`target` as it may stand in an `href` or a `src`: a runnable one
    behind a scheme no browser knows, so a reader still sees what stood
    there and nothing happens with it."""
    return f"unsafe:{target}" if runnable(target) else target


class SafeTargets:
    """Every `href` and `src` of the page through `safe_target`.

    The check sits in the parsed tree rather than in the Markdown,
    because the same scheme can be written in forms no pattern over the
    text catches: a newline inside the word, a letter as an entity, the
    colon as an entity. Measured 2026-09-10, all of those reach the tree
    as the one attribute value they became, so one check covers them.
    python-markdown asks a tree processor for `run` alone, so this needs
    no subclass and `markdown` stays an import inside the function
    below."""

    def run(self, root) -> None:
        for element in root.iter():
            for name in ("href", "src"):
                target = element.get(name)
                if target:
                    element.set(name, safe_target(target))


def html(title: str, markdown_text: str, template: str = "") -> str:
    """A complete page from the note's Markdown. `template` is the file
    `pdf_template` names: an `.html` page with `{{content}}` and
    `{{title}}` in it replaces this page, anything else is read as CSS and
    follows STYLE, so a few overriding rules are the short way in."""
    import markdown

    # The Markdown carries a video title and a channel name a stranger
    # uploaded, the URLs read out of that video's description, and a
    # model's answer written from all of those. Chrome renders this page
    # from a file:// URL to make the PDF and leaves it on disk next to it.
    # So two ways in are closed here. Raw HTML is escaped rather than
    # passed through: python-markdown lets it through by default and
    # dropped its safe mode in 3.0, and deregistering the two handlers is
    # what replaced it. A link is Markdown's own syntax, so its target is
    # checked in the tree instead. Tables, fenced code, emphasis and
    # ordinary links are untouched.
    md = markdown.Markdown(extensions=["tables", "fenced_code"])
    md.preprocessors.deregister("html_block")
    md.inlinePatterns.deregister("html")
    # Below python-markdown's own `unescape` processor, which sits at 0, so
    # every target arrives with its escaped characters already restored.
    md.treeprocessors.register(SafeTargets(), "safe_targets", -1)
    body = md.convert(markdown_text)
    style = STYLE
    if template:
        path = Path(template).expanduser()
        if not path.is_file():
            raise FetchError(f"pdf_template {path} does not exist")
        text = path.read_text(encoding="utf-8")
        if path.suffix.lower() == ".html":
            return text.replace("{{content}}", body).replace("{{title}}", escape(title))
        style = STYLE + text
    return (
        "<!doctype html>\n<html>\n<head>\n<meta charset='utf-8'>\n"
        f"<title>{escape(title)}</title>\n<style>{style}</style>\n</head>\n"
        f"<body>\n{body}\n</body>\n</html>\n"
    )


# What an install step starts with when it is meant for a shell. A step
# that reads like a sentence stays prose; the READMEs enrich reads mix
# both, and a paragraph inside a code block is unreadable.
SHELL_STARTS = ("pip", "npm", "git", "uv", "docker", "curl", "$", ">")


def is_command(step: str) -> bool:
    """Whether an install step is a line to paste into a shell: no
    sentence in it, and it opens lowercase or with a shell token."""
    text = step.strip().strip("`").strip()
    return (
        bool(text)
        and ". " not in text
        and (text[:1].islower() or text.startswith(SHELL_STARTS))
    )


def command_runs(steps: list[str]) -> list[tuple[bool, list[str]]]:
    """The steps in their order, grouped into runs of commands and runs
    of prose: `(True, [command, ...])` becomes one code block, the rest
    stays a list. Commands lose the backticks a README wraps them in."""
    runs: list[tuple[bool, list[str]]] = []
    for step in steps:
        command = is_command(step)
        text = step.strip().strip("`").strip() if command else step.strip()
        if runs and runs[-1][0] == command:
            runs[-1][1].append(text)
        else:
            runs.append((command, [text]))
    return runs


BROWSER_HINT = "`video-tldr config --set browser=<path to chrome.exe>`"


def browser(configured: str = "") -> str:
    """The Chrome or Edge to print with: the configured one, else the
    first on the PATH or in the usual places."""
    if configured:
        if Path(configured).exists() or shutil.which(configured):
            return configured
        raise FetchError(f"browser {configured} does not exist; {BROWSER_HINT}")
    for name in ("chrome", "google-chrome", "chromium", "msedge"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in BROWSERS:
        path = Path(os.path.expanduser(candidate))
        if path.exists():
            return str(path)
    raise FetchError(f"no Chrome or Edge found for the PDF; {BROWSER_HINT}")


def chromium(
    arguments: list[str], page: Path, configured_browser: str = "", task: str = ""
) -> str:
    """Run the browser headless on `page` with `arguments`; its stdout.
    Without the time budget Chrome acts before scripts ran."""
    command = [
        browser(configured_browser),
        "--headless",
        "--disable-gpu",
        "--virtual-time-budget=10000",
        *arguments,
        page.as_uri(),
    ]
    try:
        run = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=PRINT_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired as error:
        raise FetchError(
            f"{task}: no answer within {PRINT_TIMEOUT_SECONDS}s"
        ) from error
    if run.returncode != 0:
        raise FetchError(f"{task}: the browser failed: {run.stderr.strip()[:300]}")
    return run.stdout


def pdf(page: str, target: Path, configured_browser: str = "") -> Path:
    """Print `page` to `target`; the HTML stays next to it. Without the
    header flag every page carries the date and the file path."""
    source = target.with_suffix(".html")
    source.write_text(page, encoding="utf-8")
    chromium(
        ["--no-pdf-header-footer", f"--print-to-pdf={target}"],
        source,
        configured_browser,
        "pdf",
    )
    if not target.exists():
        raise FetchError("pdf: the browser printed no file")
    return target


def mermaid_page(source: str) -> str:
    """A page that draws one Mermaid diagram with the shipped script."""
    return (
        "<!doctype html>\n<html>\n<head>\n<meta charset='utf-8'>\n"
        f"<script src='{MERMAID_JS.as_uri()}'></script>\n"
        "<style>body { margin: 0; background: white; }"
        " .mermaid { display: inline-block; padding: 8px; }</style>\n</head>\n"
        f"<body>\n<pre class='mermaid'>{escape(source)}</pre>\n"
        "<script>mermaid.initialize({startOnLoad: true, theme: 'neutral'});"
        "</script>\n</body>\n</html>\n"
    )


def content_box(
    raw: bytes, width: int, height: int, margin: int = DIAGRAM_MARGIN
) -> tuple[int, int, int, int]:
    """(x, y, w, h) around what is not white in a grey frame, plus the
    margin; the whole frame when nothing is drawn."""
    import numpy as np

    # A short buffer is a screenshot that came out smaller than the window
    # asked for; reshape would answer that with a bare ValueError.
    if len(raw) < width * height:
        raise FetchError(
            f"diagram: the screenshot holds {len(raw)} bytes, "
            f"too few for {width}x{height} grey pixels"
        )
    pixels = np.frombuffer(raw[: width * height], dtype=np.uint8)
    pixels = pixels.reshape(height, width)
    rows = np.flatnonzero((pixels < 250).any(axis=1))
    cols = np.flatnonzero((pixels < 250).any(axis=0))
    if not len(rows):
        return (0, 0, width, height)
    x0, y0 = max(0, cols[0] - margin), max(0, rows[0] - margin)
    x1, y1 = min(width, cols[-1] + 1 + margin), min(height, rows[-1] + 1 + margin)
    return (int(x0), int(y0), int(x1 - x0), int(y1 - y0))


def mermaid_png(source: str, target: Path, configured_browser: str = "") -> Path:
    """Draw the diagram to `target`. The DOM says whether Mermaid could
    read the source; a screenshot of a large window is then cut to the
    drawing, because the browser's own window sizing left white space
    around the SVG on 2026-09-10."""
    page = target.with_suffix(".html")
    page.write_text(mermaid_page(source), encoding="utf-8")
    try:
        dom = chromium(["--dump-dom"], page, configured_browser, "diagram")
        if 'aria-roledescription="error"' in dom or "<svg" not in dom:
            raise FetchError("diagram: Mermaid could not read the source")
        width, height = DIAGRAM_WINDOW
        shot = target.with_name(f"{target.stem}-shot.png")
        chromium(
            [
                "--force-device-scale-factor=2",
                f"--window-size={width},{height}",
                f"--screenshot={shot}",
            ],
            page,
            configured_browser,
            "diagram",
        )
        if not shot.exists():
            raise FetchError("diagram: the browser wrote no screenshot")
        try:
            raw = ffmpeg("-i", str(shot), "-vf", "format=gray", "-f", "rawvideo", "-")
            x, y, w, h = content_box(raw, 2 * width, 2 * height)
            ffmpeg("-y", "-i", str(shot), "-vf", f"crop={w}:{h}:{x}:{y}", str(target))
        finally:
            shot.unlink(missing_ok=True)
    finally:
        page.unlink(missing_ok=True)
    return target


def code_line(document, text: str) -> None:
    """One line of a command block: Consolas on a light background.
    python-docx has no shading API, so `w:shd` goes into the paragraph by
    hand; that element is what Word paints a code block with."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run(text).font.name = "Consolas"
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    paragraph._p.get_or_add_pPr().append(shading)


def docx(
    fetched: dict,
    analysis: dict,
    placed: list[list[dict]],
    repositories: list[dict],
    labels: dict,
    folder: Path,
    target: Path,
    diagram: dict | None = None,
    timestamps: bool = True,
) -> Path:
    """The Word file. `placed` holds the pictures per section, the extra
    list at the end the ones before the first section, as `by_section`
    in render.py hands them over; the files lie in `folder`. `diagram`
    is the drawn one, placed after the summary. Without `timestamps` the
    headings, key points and captions carry the text alone."""
    from docx import Document
    from docx.shared import Inches

    def moment(seconds: float, text: str) -> str:
        return f"[{stamp(seconds)}] {text}" if timestamps else text

    vid = fetched["id"]
    doc = Document()
    doc.add_heading(fetched.get("title") or vid, 0)
    date = fetched.get("upload_date") or ""
    date = f"{date[:4]}-{date[4:6]}-{date[6:]}" if len(date) == 8 else date
    kind = labels["kinds"].get(analysis.get("kind"), analysis.get("kind"))
    doc.add_paragraph(
        f"{fetched.get('channel')} · {date} · {stamp(fetched.get('duration') or 0)}"
        f" · https://youtu.be/{vid} · {labels['kind']}: {kind}"
    )

    def picture(image: dict, caption: str = "") -> None:
        path = folder / image["file"]
        if not path.exists():
            return
        doc.add_picture(str(path), width=Inches(6))
        if caption:
            doc.add_paragraph().add_run(caption).italic = True
        for command in image.get("commands") or []:
            code_line(doc, command)

    if fetched.get("thumbnail"):
        picture({"file": fetched["thumbnail"]})
    doc.add_heading(labels["summary"], 1)
    doc.add_paragraph((analysis.get("summary") or "").strip())
    if diagram:
        picture(diagram, diagram.get("caption", ""))
    doc.add_heading(labels["sections"], 1)
    sections = analysis.get("sections") or []
    for image in placed[-1]:
        picture(image, moment(image["time"], image.get("caption", "")))
    for section, pictures in zip(sections, placed):
        doc.add_heading(moment(section["start"], section["title"]), 2)
        doc.add_paragraph(section["summary"].strip())
        for image in pictures:
            picture(image, moment(image["time"], image.get("caption", "")))
    if analysis.get("key_points"):
        doc.add_heading(labels["key_points"], 1)
        for point in analysis["key_points"]:
            doc.add_paragraph(moment(point["time"], point["text"]), style="List Bullet")
    if repositories:
        doc.add_heading(labels["install"], 1)
        for repository in repositories:
            doc.add_heading(f"{repository['repo']} ({repository['url']})", 2)
            if repository.get("what"):
                doc.add_paragraph(repository["what"].strip())
            for commands, steps in command_runs(repository.get("install") or []):
                for step in steps:
                    if commands:
                        code_line(doc, step)
                    else:
                        doc.add_paragraph(step.replace("`", ""), style="List Number")
    if analysis.get("links"):
        doc.add_heading(labels["links"], 1)
        for link in analysis["links"]:
            doc.add_paragraph(f"{link['url']} ({link['role']})", style="List Bullet")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target
